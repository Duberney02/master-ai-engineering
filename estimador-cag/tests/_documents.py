"""Documentos PDF y Word mínimos generados en las pruebas (sin archivos binarios en el repositorio)."""

import io

from docx import Document


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(*pages: str) -> bytes:
    """PDF con una página por texto (ASCII), con tabla de referencias cruzadas correcta."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] /Count {len(pages)} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, text in zip(page_ids, pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({_escape(text)}) Tj ET"
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {page_id + 1} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_blank_pdf() -> bytes:
    """PDF válido sin texto (como un documento escaneado)."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
